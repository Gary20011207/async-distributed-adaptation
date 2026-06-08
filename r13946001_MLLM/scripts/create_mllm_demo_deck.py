from __future__ import annotations

import textwrap
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


MLLM_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = MLLM_ROOT.parent
REPORT_DIR = MLLM_ROOT / "figures" / "report"
PRESENTATION_DIR = MLLM_ROOT / "presentation"
ASSET_DIR = PRESENTATION_DIR / "demo_assets"
SOURCE_PPTX = REPO_ROOT / "Group 4：PP-DMA_ Privacy-Preserving Distributed Model Adaptation.pptx"
OUT_PPTX = REPO_ROOT / "Group 4：PP-DMA_ Privacy-Preserving Distributed Model Adaptation_MLLM_Demo.pptx"
OUT_MD = PRESENTATION_DIR / "mllm_demo_section_zh.md"


COLORS = {
    "Sync": "#6B7280",
    "Naive Async": "#F59E0B",
    "Staleness": "#EF4444",
    "FedBuff": "#2563EB",
    "CAA-v2": "#059669",
    "Base Qwen": "#7C3AED",
}
METHOD_ORDER = ["Sync", "Naive Async", "Staleness", "FedBuff", "CAA-v2"]


def pct(x: float) -> str:
    return f"{100 * x:.1f}%"


def savefig(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()


def text_box(slide, x, y, w, h, text, size=20, bold=False, color=(31, 41, 55), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Noto Sans CJK TC"
    run.font.color.rgb = RGBColor(*color)
    if align:
        p.alignment = align
    return box


def bullet_box(slide, x, y, w, h, bullets, size=18, color=(31, 41, 55)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Noto Sans CJK TC"
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(7)
    return box


def add_title(slide, title, subtitle=None):
    text_box(slide, 0.55, 0.35, 12.1, 0.45, title, size=24, bold=True, color=(15, 23, 42))
    if subtitle:
        text_box(slide, 0.58, 0.82, 12.0, 0.28, subtitle, size=11, color=(100, 116, 139))
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.55),
        Inches(1.14),
        Inches(12.2),
        Inches(0.02),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(5, 150, 105)
    line.line.fill.background()


def add_footer(slide, note="PP-DMA DEMO | LLM/MLLM federated adaptation"):
    text_box(slide, 0.55, 7.12, 8.0, 0.2, note, size=8, color=(148, 163, 184))


def set_bg(slide, color=(248, 250, 252)):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(*color)


def add_metric_card(slide, x, y, w, h, label, value, caption, accent=(5, 150, 105)):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(226, 232, 240)
    text_box(slide, x + 0.18, y + 0.15, w - 0.36, 0.25, label, size=10, bold=True, color=accent)
    text_box(slide, x + 0.18, y + 0.43, w - 0.36, 0.38, value, size=23, bold=True, color=(15, 23, 42))
    text_box(slide, x + 0.18, y + 0.9, w - 0.36, h - 0.95, caption, size=8, color=(71, 85, 105))


def make_figures() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    mean = pd.read_csv(REPORT_DIR / "mean_std_summary.csv")
    validity = pd.read_csv(REPORT_DIR / "qwenvl_method_validity.csv")
    base_diag = pd.read_csv(REPORT_DIR / "qwenvl_closed_answer_diagnostics_summary.csv")
    delay = pd.read_csv(REPORT_DIR / "delay_stress_summary.csv")
    samples = pd.read_csv(REPORT_DIR / "qwenvl_method_prediction_samples.csv")

    assets: dict[str, Path] = {}

    text = mean[mean["task_family"].eq("text_mcqa")]
    agg = (
        text.groupby("method_label", as_index=False)[
            ["best_acc_mean", "final_acc_mean", "stability_drop_mean"]
        ]
        .mean()
        .set_index("method_label")
        .reindex(METHOD_ORDER)
        .reset_index()
    )
    fig, ax = plt.subplots(figsize=(8.4, 4.6))
    x = np.arange(len(agg))
    width = 0.36
    ax.bar(x - width / 2, agg["best_acc_mean"], width, label="Best Acc", color="#0F766E")
    ax.bar(x + width / 2, agg["final_acc_mean"], width, label="Final Acc", color="#60A5FA")
    for i, row in agg.iterrows():
        ax.text(i - width / 2, row.best_acc_mean + 0.008, pct(row.best_acc_mean), ha="center", fontsize=8)
        ax.text(i + width / 2, row.final_acc_mean + 0.008, pct(row.final_acc_mean), ha="center", fontsize=8)
    ax.set_ylim(0.24, 0.45)
    ax.set_ylabel("Accuracy")
    ax.set_title("Text MCQA: mean over MMLU, MedMCQA, MedQA-USMLE, PubMedQA")
    ax.set_xticks(x)
    ax.set_xticklabels(agg["method_label"], rotation=18, ha="right")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, ncol=2, loc="upper left")
    assets["text_mcqa"] = ASSET_DIR / "demo_text_mcqa_best_final.png"
    savefig(assets["text_mcqa"])

    real = mean[mean["backend_group"].eq("real_qwenvl")].copy()
    pivot = real.pivot_table(index="dataset", columns="method_label", values="final_acc_mean", aggfunc="mean")
    pivot = pivot.reindex(columns=METHOD_ORDER)
    fig, ax = plt.subplots(figsize=(8.4, 4.3))
    x = np.arange(len(pivot.index))
    width = 0.15
    offsets = np.linspace(-2, 2, len(METHOD_ORDER)) * width
    for method, off in zip(METHOD_ORDER, offsets):
        vals = pivot[method].values
        ax.bar(x + off, vals, width, label=method, color=COLORS[method])
        for xi, v in zip(x + off, vals):
            ax.text(xi, v + 0.012, pct(v), ha="center", fontsize=7, rotation=90)
    ax.set_ylim(0.45, 0.94)
    ax.set_ylabel("Final accuracy")
    ax.set_title("Real Qwen2.5-VL 3B 4-bit QLoRA: small fair-budget VQA")
    ax.set_xticks(x)
    ax.set_xticklabels(["PathVQA", "VQA-RAD"])
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, ncol=5, loc="upper center", bbox_to_anchor=(0.5, -0.12), fontsize=8)
    assets["real_qwenvl"] = ASSET_DIR / "demo_real_qwenvl_final_acc.png"
    savefig(assets["real_qwenvl"])

    base = base_diag[["dataset", "accuracy"]].copy()
    base["method_label"] = "Base Qwen"
    comp = validity[validity["method_label"].isin(["CAA-v2", "Naive Async", "Sync"])][
        ["dataset", "method_label", "accuracy"]
    ]
    before_after = pd.concat([base, comp], ignore_index=True)
    order = ["Base Qwen", "Sync", "Naive Async", "CAA-v2"]
    pivot = before_after.pivot_table(index="dataset", columns="method_label", values="accuracy", aggfunc="mean")
    pivot = pivot.reindex(columns=order)
    fig, ax = plt.subplots(figsize=(8.4, 4.2))
    x = np.arange(len(pivot.index))
    width = 0.18
    offsets = np.linspace(-1.5, 1.5, len(order)) * width
    for method, off in zip(order, offsets):
        vals = pivot[method].values
        ax.bar(x + off, vals, width, label=method, color=COLORS.get(method, "#94A3B8"))
        for xi, v in zip(x + off, vals):
            ax.text(xi, v + 0.012, pct(v), ha="center", fontsize=8, rotation=90)
    ax.set_ylim(0.38, 0.98)
    ax.set_ylabel("Diagnostic accuracy")
    ax.set_title("Closed-answer diagnostic probe: base model vs FL checkpoints")
    ax.set_xticks(x)
    ax.set_xticklabels(["PathVQA", "VQA-RAD"])
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.12), fontsize=8)
    assets["before_after"] = ASSET_DIR / "demo_base_vs_fl_diagnostic.png"
    savefig(assets["before_after"])

    vqa_delay = delay[delay["dataset"].eq("vqa_rad_closed")].copy()
    pivot = vqa_delay.pivot_table(index="delay_mode", columns="method_label", values="final_acc", aggfunc="mean")
    pivot = pivot.reindex(index=["uniform", "lognormal"], columns=["Naive Async", "Staleness", "FedBuff", "CAA-v2"])
    fig, ax = plt.subplots(figsize=(8.4, 4.2))
    x = np.arange(len(pivot.index))
    width = 0.18
    order2 = ["Naive Async", "Staleness", "FedBuff", "CAA-v2"]
    offsets = np.linspace(-1.5, 1.5, len(order2)) * width
    for method, off in zip(order2, offsets):
        vals = pivot[method].values
        ax.bar(x + off, vals, width, label=method, color=COLORS[method])
        for xi, v in zip(x + off, vals):
            ax.text(xi, v + 0.008, pct(v), ha="center", fontsize=8, rotation=90)
    ax.set_ylim(0.64, 0.78)
    ax.set_ylabel("Final accuracy")
    ax.set_title("VQA-RAD proxy delay stress: final accuracy under async timing")
    ax.set_xticks(x)
    ax.set_xticklabels(["Uniform delay", "Lognormal delay"])
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.12), fontsize=8)
    assets["delay"] = ASSET_DIR / "demo_vqa_rad_delay_stress.png"
    savefig(assets["delay"])

    table_rows = []
    preferred = samples[
        (samples["dataset"].eq("vqa_rad_closed"))
        & (samples["method_label"].isin(["CAA-v2", "Naive Async", "Sync"]))
    ].copy()
    selected_rows = []
    sample_plan = {"Sync": [0, 1], "Naive Async": [1, 2], "CAA-v2": [2, 3]}
    for method, indices in sample_plan.items():
        group = preferred[preferred["method_label"].eq(method)].sort_values(["sample_index", "seed"])
        for sample_index in indices:
            match = group[group["sample_index"].eq(sample_index)]
            if not match.empty:
                selected_rows.append(match.iloc[0])
    preferred = pd.DataFrame(selected_rows)
    for _, row in preferred.iterrows():
        q = textwrap.shorten(str(row["question"]), width=54, placeholder="...")
        table_rows.append(
            [
                row["method_label"],
                q,
                str(row["parsed_answer"]),
                str(row["gold_answer"]),
                "OK" if bool(row["correct"]) else "Miss",
            ]
        )
    fig, ax = plt.subplots(figsize=(10.5, 3.9))
    ax.axis("off")
    tbl = ax.table(
        cellText=table_rows,
        colLabels=["Method", "Question", "Pred.", "Gold", "Result"],
        loc="center",
        cellLoc="left",
        colWidths=[0.13, 0.54, 0.09, 0.09, 0.1],
    )
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(8.5)
    tbl.scale(1, 1.65)
    for (r, c), cell in tbl.get_celld().items():
        cell.set_edgecolor("#CBD5E1")
        if r == 0:
            cell.set_facecolor("#E2E8F0")
            cell.set_text_props(weight="bold")
        elif c == 4:
            cell.set_facecolor("#DCFCE7" if cell.get_text().get_text() == "OK" else "#FEE2E2")
    ax.set_title("Real Qwen2.5-VL checkpoint prediction samples", pad=12)
    assets["samples"] = ASSET_DIR / "demo_qwenvl_prediction_samples.png"
    savefig(assets["samples"])

    return assets


def add_architecture(slide):
    labels = [
        ("Hospital A\nprivate QA/VQA", 0.8, 2.1),
        ("Hospital B\nprivate QA/VQA", 0.8, 3.35),
        ("Hospital C\nprivate QA/VQA", 0.8, 4.6),
        ("LoRA adapter\nupdate only", 3.55, 3.25),
        ("Clockless server\nCAA-v2 aggregation", 6.0, 3.25),
        ("Global adapter\nclosed answer demo", 9.05, 3.25),
    ]
    boxes = []
    for text, x, y in labels:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.1), Inches(0.75)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(148, 163, 184)
        text_box(slide, x + 0.1, y + 0.14, 1.9, 0.45, text, size=11, bold=True, color=(15, 23, 42), align=PP_ALIGN.CENTER)
        boxes.append(shape)
    arrow_pairs = [
        ((2.9, 2.48), (3.55, 3.63)),
        ((2.9, 3.73), (3.55, 3.63)),
        ((2.9, 4.98), (3.55, 3.63)),
        ((5.65, 3.63), (6.0, 3.63)),
        ((8.1, 3.63), (9.05, 3.63)),
    ]
    for (x1, y1), (x2, y2) in arrow_pairs:
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = RGBColor(5, 150, 105)
        line.line.width = Pt(2)
        line.line.end_arrowhead = True
    text_box(slide, 3.45, 4.35, 5.0, 0.55, "server 不等待所有 client；只用 logical version 量測 staleness", size=12, color=(71, 85, 105), align=PP_ALIGN.CENTER)


def add_img(slide, path: Path, x, y, w, h=None):
    if h is None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    else:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_demo_slides(assets: dict[str, Path]) -> None:
    prs = Presentation(SOURCE_PPTX)
    blank_layout = prs.slide_layouts[6]

    def new_slide(title, subtitle=None):
        slide = prs.slides.add_slide(blank_layout)
        set_bg(slide)
        add_title(slide, title, subtitle)
        add_footer(slide)
        return slide

    slide = new_slide("08 DEMO EXTENSION｜Medical LLM / MLLM Federated Adaptation")
    text_box(slide, 0.75, 1.55, 11.8, 0.8, "從醫學影像分類，延伸到 closed-ended medical QA / VQA", size=26, bold=True, color=(15, 23, 42), align=PP_ALIGN.CENTER)
    bullet_box(
        slide,
        1.05,
        2.65,
        11.0,
        2.5,
        [
            "DEMO 目標：比較 base model 與 FL adapter checkpoint 的 closed-answer 行為。",
            "訓練形式：client 只更新 LoRA / QLoRA adapter，不上傳原始醫療文字或影像。",
            "分散式重點：保留 no global clock、staleness、straggler、fair update budget。",
            "報告定位：text MCQA 是主要量化矩陣；real Qwen2.5-VL 是小規模 feasibility demo。",
        ],
        size=18,
    )
    add_metric_card(slide, 1.05, 5.55, 2.55, 1.05, "official summaries", "160", "所有正式 MLLM/LLM summary rows")
    add_metric_card(slide, 3.9, 5.55, 2.55, 1.05, "primary rows", "121", "去重後 IID fair-budget headline rows")
    add_metric_card(slide, 6.75, 5.55, 2.55, 1.05, "real Qwen2.5-VL", "30", "4-bit QLoRA 小矩陣 rows")
    add_metric_card(slide, 9.6, 5.55, 2.55, 1.05, "prediction probes", "80", "checkpoint closed-answer samples")

    slide = new_slide("LLM/MLLM DEMO 架構", "用同一個 clockless async FL 框架，改成聚合 adapter weights。")
    add_architecture(slide)
    bullet_box(
        slide,
        0.85,
        5.35,
        11.8,
        1.2,
        [
            "Raw medical data stays local：文字題目、影像、答案都不離開 client。",
            "Server receives adapter deltas：比較 Sync / Naive / Staleness / FedBuff / CAA-v2。",
            "Evaluation parser forces finite answer space：A/B/C/D 或 yes/no，並記錄 invalid answer rate。",
        ],
        size=15,
    )

    slide = new_slide("實驗矩陣與公平性", "DEMO 結果不混淆 proxy 與 real Qwen2.5-VL。")
    bullet_box(
        slide,
        0.8,
        1.45,
        5.55,
        3.9,
        [
            "Text MCQA：MMLU、MedMCQA、MedQA-USMLE、PubMedQA。",
            "Proxy VQA：VQA-RAD、PathVQA，用來做算法與 pipeline 大量驗證。",
            "Real MLLM：Qwen2.5-VL-3B-Instruct 4-bit QLoRA，小樣本 fair-budget feasibility。",
            "Fair budget：Async events = Sync rounds × clients。",
            "所有 headline group：same seed、same local epochs、same delay mode、same client count。",
        ],
        size=17,
    )
    add_metric_card(slide, 7.0, 1.55, 2.35, 1.15, "fair groups", "24/24", "fairness audit 全部通過")
    add_metric_card(slide, 9.75, 1.55, 2.35, 1.15, "text seeds", "3", "每個 text dataset 三個 seed")
    add_metric_card(slide, 7.0, 3.05, 2.35, 1.15, "delay stress", "16", "獨立報告，不放入 headline mean")
    add_metric_card(slide, 9.75, 3.05, 2.35, 1.15, "non-IID", "10", "Dirichlet hospital scenario")
    text_box(slide, 7.0, 4.85, 5.25, 1.1, "重要：real Qwen2.5-VL 的數量級仍是 feasibility，不把它宣稱成大規模 MLLM SOTA。", size=16, bold=True, color=(190, 18, 60))

    slide = new_slide("Text MCQA 結果：CAA-v2 提高 peak，Sync final 較穩", "Qwen1.5-0.5B LoRA；四個 closed-ended text QA datasets 平均。")
    add_img(slide, assets["text_mcqa"], 0.75, 1.45, 7.4)
    bullet_box(
        slide,
        8.55,
        1.62,
        4.05,
        4.5,
        [
            "CAA-v2 best mean：38.9%，高於 FedBuff / Naive / Staleness / Sync。",
            "Sync final mean：35.4%，CAA-v2 final mean：34.9%，代表 async peak 有提升但 final stability 還有空間。",
            "Demo 重點：FL adapter 可以在有限題型上被公平比較，不只是單次 prompt 測試。",
        ],
        size=15,
    )

    slide = new_slide("Real Qwen2.5-VL：小規模 fair-budget VQA feasibility", "Qwen2.5-VL-3B-Instruct 4-bit QLoRA，只聚合 adapter。")
    add_img(slide, assets["real_qwenvl"], 0.75, 1.45, 7.45)
    bullet_box(
        slide,
        8.55,
        1.6,
        4.1,
        4.75,
        [
            "VQA-RAD：CAA-v2 / FedBuff / Sync final 都是 83.3%。Naive/Staleness peak 87.5%，但 final 回到 83.3%。",
            "PathVQA：Sync final 66.7%，async family 約 62.5%。",
            "目前結論：real MLLM pipeline 可跑通；結果是 feasibility，不是充分統計勝出。",
        ],
        size=15,
    )

    slide = new_slide("訓練前後診斷：要展示，也要誠實", "8-sample closed-answer probe；用來 demo 行為，不作為 headline benchmark。")
    add_img(slide, assets["before_after"], 0.75, 1.45, 7.45)
    bullet_box(
        slide,
        8.55,
        1.62,
        4.1,
        4.8,
        [
            "Base Qwen2.5-VL 解析率 100%，兩個小樣本 accuracy 都是 87.5%。",
            "FL checkpoints 解析率也維持 100%，代表 adapter 不會破壞 closed-answer format。",
            "但 FL 後不一定比 base 更高：這提醒 demo 要說 proof-of-feasibility，而不是過度宣稱。",
            "真正要做更強 before/after，需要固定同一批 eval questions 並擴大 sample size。",
        ],
        size=14,
    )

    slide = new_slide("Closed-ended VQA Prediction Samples", "展示 FL checkpoint 實際輸出、parser、gold answer 與 correct/miss。")
    add_img(slide, assets["samples"], 0.55, 1.45, 12.2)
    text_box(slide, 0.8, 6.45, 11.7, 0.45, "DEMO 可現場呈現：輸入 medical image + question，模型只需回答有限答案空間；parser 將自由文字轉成 A/B/yes/no。", size=16, bold=True, color=(15, 23, 42), align=PP_ALIGN.CENTER)

    slide = new_slide("分散式壓力測試：delay distribution 會改變結果", "VQA-RAD proxy stress；stress results 獨立於 headline mean±std。")
    add_img(slide, assets["delay"], 0.75, 1.45, 7.45)
    bullet_box(
        slide,
        8.55,
        1.62,
        4.1,
        4.9,
        [
            "Naive/Staleness 的 avg staleness 約 8；FedBuff/CAA-v2 約 1.5。",
            "Uniform delay 下 CAA-v2 final 75.0%，高於 FedBuff 73.3%、Naive 70.3%。",
            "Lognormal delay 下 CAA-v2 final 74.3%，仍高於其他 async baselines。",
            "這是分散式系統故事：不是只看 accuracy，而是看 delay、staleness、client timing。",
        ],
        size=14,
    )

    slide = new_slide("LLM/MLLM DEMO Takeaways", "這一章補足 PP-DMA 從 CNN biomedical FL 延伸到 medical QA/VQA 的可行性。")
    bullet_box(
        slide,
        0.9,
        1.55,
        11.6,
        3.5,
        [
            "已完成：text MCQA 三種 seed 的 fair-budget FL 矩陣，real Qwen2.5-VL 4-bit QLoRA 小矩陣，以及 80 筆 checkpoint prediction diagnostics。",
            "可展示：base model vs FL checkpoint 的 closed-answer behavior、invalid answer rate、best/final accuracy、stability drop。",
            "保守結論：CAA-v2 在 text MCQA peak 與 VQA delay stress 上有優勢；real MLLM 目前證明 pipeline 可行，還不是 SOTA claim。",
            "下一步 DEMO：現場固定幾題 medical QA/VQA，切換 Base / Sync / CAA-v2 adapter，展示 answer、parser、correctness 與 latency。",
        ],
        size=19,
    )
    add_metric_card(slide, 1.2, 5.55, 3.1, 1.05, "Main evidence", "Text MCQA", "完整 multi-seed quantitative matrix")
    add_metric_card(slide, 5.05, 5.55, 3.1, 1.05, "Feasibility", "Real Qwen2.5-VL", "3 seeds × 2 VQA datasets × 5 methods")
    add_metric_card(slide, 8.9, 5.55, 3.1, 1.05, "Demo metric", "Valid + Correct", "finite answer parsing and accuracy")

    prs.save(OUT_PPTX)


def write_md() -> None:
    text = """# LLM / MLLM DEMO 章節講稿備份

## 08 DEMO Extension

這一段是把原本 MedMNIST 的 CNN 影像分類實驗，延伸到 LLM / MLLM 的 closed-ended medical QA / VQA。Demo 的核心不是宣稱模型已經達到醫療 SOTA，而是展示同一個 clockless asynchronous FL 框架，也能用在 LoRA / QLoRA adapter adaptation。

## LLM/MLLM Demo 架構

每個 hospital client 保留自己的題目、影像和答案，只在本地更新 adapter。Server 不等待所有 client，而是依照 update arrival event 聚合 adapter delta。答案空間被限制在 A/B/C/D 或 yes/no，所以可以明確計算 valid answer rate 和 accuracy。

## 實驗矩陣

目前官方 summary 有 160 rows，headline 使用 121 個去重後 IID fair-budget rows。Text MCQA 使用 MMLU、MedMCQA、MedQA-USMLE、PubMedQA；VQA 包含 VQA-RAD 和 PathVQA。Proxy VQA 用來做完整算法驗證，real Qwen2.5-VL 則是 4-bit QLoRA feasibility。

## Text MCQA Result

在四個 text MCQA datasets 平均下，CAA-v2 的 best accuracy mean 約 38.9%，高於其他方法；Sync final accuracy mean 約 35.4%，CAA-v2 約 34.9%。這代表 CAA-v2 的 peak 有提升，但 async final stability 還有改善空間。

## Real Qwen2.5-VL Result

Real Qwen2.5-VL 的 VQA-RAD 小矩陣中，CAA-v2、FedBuff、Sync 的 final accuracy 都是 83.3%；Naive 和 Staleness peak 可以到 87.5%，但 final 回到 83.3%。PathVQA 上 Sync 約 66.7%，async family 約 62.5%。這是 feasibility evidence，不是大規模 MLLM SOTA claim。

## Before / After Diagnostic

Base Qwen2.5-VL 的小樣本 closed-answer diagnostic valid rate 是 100%，accuracy 是 87.5%。FL checkpoints 也維持 100% valid rate，表示 adapter 不會破壞 closed-answer format。不過 FL 後不一定比 base 更高，所以這部分要誠實說是 demo 診斷，不是嚴格勝出證明。

## Distributed Stress

在 VQA-RAD proxy delay stress 中，CAA-v2 在 uniform 和 lognormal delay 下都有較高 final accuracy。這支持我們的分散式系統觀點：async FL 不只是 ML optimizer，也受 delay distribution、staleness 和 client timing 影響。
"""
    OUT_MD.parent.mkdir(parents=True, exist_ok=True)
    OUT_MD.write_text(text, encoding="utf-8")


def main() -> None:
    if not SOURCE_PPTX.exists():
        raise FileNotFoundError(SOURCE_PPTX)
    assets = make_figures()
    add_demo_slides(assets)
    write_md()
    print(f"wrote {OUT_PPTX}")
    print(f"wrote {OUT_MD}")
    print("assets:")
    for path in assets.values():
        print(path)


if __name__ == "__main__":
    main()
